from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

PDF_EXTENSION = ".pdf"
INDEX_EXTENSIONS = {".xlsx", ".xlsm", ".csv"}
EMPTY_VALUES = {"", "none", "null", "nan", "n/a", "na", "-"}
PERSONNUMMER_HEADERS = {
    "personnummer",
    "personnr",
    "personnr.",
    "person number",
    "personal identity number",
    "person_id",
}


@dataclass
class PdfResult:
    input_pdf: Path
    expected_output_dir: Path
    split_files: list[Path] = field(default_factory=list)
    missing_numbers: list[int] = field(default_factory=list)

    @property
    def successful(self) -> bool:
        return bool(self.split_files)


@dataclass
class VolumeResult:
    relative_dir: Path
    output_dir_exists: bool
    total_pdfs: int
    successful_pdfs: int
    failed_pdfs: int

    @property
    def successful(self) -> bool:
        return self.output_dir_exists and self.total_pdfs > 0 and self.failed_pdfs == 0


@dataclass
class IndexResult:
    index_file: Path | None = None
    sheet_name: str | None = None
    personnummer_column: str | None = None
    total_rows: int = 0
    extracted: int = 0
    missing: int = 0
    missing_rows: list[str] = field(default_factory=list)
    error: str | None = None


def natural_key(value: Path | str) -> list[Any]:
    parts = re.split(r"(\d+)", str(value).lower())
    return [int(part) if part.isdigit() else part for part in parts]


def normalized_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() in EMPTY_VALUES else text


def normalized_header(value: Any) -> str:
    return re.sub(r"\s+", " ", normalized_text(value).lower().replace("_", " "))


def find_input_pdfs(input_root: Path) -> list[Path]:
    return sorted(
        (path for path in input_root.rglob("*") if path.is_file() and path.suffix.lower() == PDF_EXTENSION),
        key=natural_key,
    )


def find_split_files(original_pdf: Path, output_dir: Path) -> tuple[list[Path], list[int]]:
    if not output_dir.is_dir():
        return [], []

    # The suffix is checked after the complete original stem. Therefore an
    # original ending in "_" matches output files ending in "__1", "__2", etc.
    pattern = re.compile(rf"^{re.escape(original_pdf.stem)}_(\d+)$", re.IGNORECASE)
    numbered_files: list[tuple[int, Path]] = []

    for candidate in output_dir.iterdir():
        if not candidate.is_file() or candidate.suffix.lower() != PDF_EXTENSION:
            continue
        match = pattern.fullmatch(candidate.stem)
        if match:
            numbered_files.append((int(match.group(1)), candidate))

    numbered_files.sort(key=lambda item: (item[0], natural_key(item[1])))
    numbers = sorted({number for number, _ in numbered_files})
    missing_numbers = []
    if numbers:
        missing_numbers = sorted(set(range(1, max(numbers) + 1)) - set(numbers))

    return [path for _, path in numbered_files], missing_numbers


def validate_pdfs(input_root: Path, output_root: Path) -> list[PdfResult]:
    results: list[PdfResult] = []
    for input_pdf in find_input_pdfs(input_root):
        relative_parent = input_pdf.parent.relative_to(input_root)
        output_dir = output_root / relative_parent
        split_files, missing_numbers = find_split_files(input_pdf, output_dir)
        results.append(PdfResult(input_pdf, output_dir, split_files, missing_numbers))
    return results


def validate_volumes(input_root: Path, output_root: Path, pdf_results: list[PdfResult]) -> list[VolumeResult]:
    grouped: dict[Path, list[PdfResult]] = {}
    for result in pdf_results:
        grouped.setdefault(result.input_pdf.parent, []).append(result)

    volume_results: list[VolumeResult] = []
    for input_dir in sorted(grouped, key=natural_key):
        relative_dir = input_dir.relative_to(input_root)
        output_dir = output_root / relative_dir
        entries = grouped[input_dir]
        successful = sum(result.successful for result in entries)
        volume_results.append(
            VolumeResult(
                relative_dir=relative_dir,
                output_dir_exists=output_dir.is_dir(),
                total_pdfs=len(entries),
                successful_pdfs=successful,
                failed_pdfs=len(entries) - successful,
            )
        )
    return volume_results


def locate_index_file(output_root: Path, selected: Path | None) -> tuple[Path | None, str | None]:
    if selected is not None:
        return selected, None

    candidates = sorted(
        (path for path in output_root.rglob("*") if path.is_file() and path.suffix.lower() in INDEX_EXTENSIONS),
        key=natural_key,
    )
    named_index = [path for path in candidates if "index" in path.stem.lower()]

    if len(named_index) == 1:
        return named_index[0], None
    if len(candidates) == 1:
        return candidates[0], None
    if not candidates:
        return None, "Ingen indexfil hittades. Ange den med --index."
    return None, "Flera möjliga indexfiler hittades. Ange rätt fil med --index."


def find_personnummer_column(headers: list[Any], requested: str | None) -> int | None:
    normalized = [normalized_header(header) for header in headers]
    if requested:
        target = normalized_header(requested)
        return next((i for i, header in enumerate(normalized) if header == target), None)
    valid_headers = {normalized_header(header) for header in PERSONNUMMER_HEADERS}
    return next((i for i, header in enumerate(normalized) if header in valid_headers), None)


def describe_row(headers: list[Any], row: list[Any], row_number: int) -> str:
    preferred = {"filnamn", "filename", "pdf", "volym", "volume", "efternamn", "förnamn", "fornamn", "namn", "name"}
    details: list[str] = []
    for index, header in enumerate(headers):
        if normalized_header(header) in preferred and index < len(row):
            value = normalized_text(row[index])
            if value:
                details.append(f"{normalized_text(header) or f'Kolumn {index + 1}'}: {value}")
    if not details:
        for index, value in enumerate(row):
            text = normalized_text(value)
            if text:
                header = normalized_text(headers[index]) if index < len(headers) else f"Kolumn {index + 1}"
                details.append(f"{header}: {text}")
            if len(details) == 3:
                break
    return f"Rad {row_number}: " + (" | ".join(details) if details else "Ingen identifierande information")


def process_index_rows(
    index_file: Path,
    headers: list[Any],
    rows: Iterable[tuple[int, list[Any]]],
    requested_column: str | None,
    sheet_name: str | None = None,
) -> IndexResult:
    result = IndexResult(index_file=index_file, sheet_name=sheet_name)
    column_index = find_personnummer_column(headers, requested_column)
    if column_index is None:
        available = ", ".join(normalized_text(header) for header in headers if normalized_text(header))
        result.error = f"Personnummerkolumnen hittades inte. Tillgängliga kolumner: {available}"
        return result

    result.personnummer_column = normalized_text(headers[column_index])
    for row_number, row in rows:
        if not any(normalized_text(value) for value in row):
            continue
        result.total_rows += 1
        value = row[column_index] if column_index < len(row) else None
        if normalized_text(value):
            result.extracted += 1
        else:
            result.missing += 1
            result.missing_rows.append(describe_row(headers, row, row_number))
    return result


def validate_excel_index(index_file: Path, requested_column: str | None, requested_sheet: str | None) -> IndexResult:
    if load_workbook is None:
        return IndexResult(index_file=index_file, error="openpyxl saknas. Installera med: pip install openpyxl")
    try:
        workbook = load_workbook(index_file, read_only=True, data_only=True)
        try:
            if requested_sheet:
                if requested_sheet not in workbook.sheetnames:
                    return IndexResult(index_file=index_file, error=f"Excel-bladet '{requested_sheet}' finns inte.")
                worksheet = workbook[requested_sheet]
            else:
                worksheet = workbook.active

            iterator = worksheet.iter_rows(values_only=True)
            first_row = next(iterator, None)
            if first_row is None:
                return IndexResult(index_file=index_file, sheet_name=worksheet.title, error="Indexfilen är tom.")
            headers = list(first_row)
            rows = ((number, list(row)) for number, row in enumerate(iterator, start=2))
            return process_index_rows(index_file, headers, rows, requested_column, worksheet.title)
        finally:
            workbook.close()
    except Exception as exc:
        return IndexResult(index_file=index_file, error=f"Kunde inte läsa Excel-filen: {exc}")


def read_csv(index_file: Path) -> list[list[str]]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            with index_file.open("r", encoding=encoding, newline="") as handle:
                sample = handle.read(8192)
                handle.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                    return list(csv.reader(handle, dialect))
                except csv.Error:
                    handle.seek(0)
                    return list(csv.reader(handle, delimiter=";"))
        except UnicodeDecodeError as exc:
            last_error = exc
    raise RuntimeError(f"CSV-filen kunde inte avkodas: {last_error}")


def validate_csv_index(index_file: Path, requested_column: str | None) -> IndexResult:
    try:
        all_rows = read_csv(index_file)
    except Exception as exc:
        return IndexResult(index_file=index_file, error=f"Kunde inte läsa CSV-filen: {exc}")
    if not all_rows:
        return IndexResult(index_file=index_file, error="Indexfilen är tom.")
    headers = list(all_rows[0])
    rows = ((number, list(row)) for number, row in enumerate(all_rows[1:], start=2))
    return process_index_rows(index_file, headers, rows, requested_column)


def validate_index(index_file: Path | None, discovery_error: str | None, requested_column: str | None, requested_sheet: str | None) -> IndexResult:
    if discovery_error:
        return IndexResult(index_file=index_file, error=discovery_error)
    if index_file is None or not index_file.is_file():
        return IndexResult(index_file=index_file, error=f"Indexfilen finns inte: {index_file}")
    suffix = index_file.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return validate_excel_index(index_file, requested_column, requested_sheet)
    if suffix == ".csv":
        return validate_csv_index(index_file, requested_column)
    return IndexResult(index_file=index_file, error=f"Indexformatet stöds inte: {suffix}")


def create_report(
    report_file: Path,
    input_root: Path,
    output_root: Path,
    pdf_results: list[PdfResult],
    volume_results: list[VolumeResult],
    index_result: IndexResult,
) -> None:
    successful_pdfs = sum(result.successful for result in pdf_results)
    failed_pdfs = len(pdf_results) - successful_pdfs
    successful_volumes = sum(result.successful for result in volume_results)
    failed_volumes = len(volume_results) - successful_volumes
    split_file_count = sum(len(result.split_files) for result in pdf_results)

    lines = [
        "=" * 78,
        "VALIDERINGSRAPPORT FÖR EXAMENSBEVIS",
        "=" * 78,
        f"Skapad: {datetime.now().astimezone().strftime('%Y-%m-%d %H:%M:%S %Z')}",
        f"Indata: {input_root}",
        f"Utdata: {output_root}",
        f"Index: {index_result.index_file or 'Ingen identifierad'}",
        "",
        "SAMMANFATTNING",
        "-" * 78,
        f"Volymer totalt: {len(volume_results)}",
        f"Volymer fullständigt bearbetade: {successful_volumes}",
        f"Volymer saknade eller ofullständiga: {failed_volumes}",
        f"Ursprungliga PDF-filer: {len(pdf_results)}",
        f"PDF-filer med numrerad utdata: {successful_pdfs}",
        f"PDF-filer utan numrerad utdata: {failed_pdfs}",
        f"Numrerade del-PDF-filer som hittades: {split_file_count}",
    ]

    if index_result.error:
        lines.extend(["Indexkontroll: MISSLYCKADES", f"Orsak: {index_result.error}"])
    else:
        lines.extend([
            f"Indexrader: {index_result.total_rows}",
            f"Personnummer extraherade: {index_result.extracted}",
            f"Personnummer saknas: {index_result.missing}",
        ])

    lines.extend(["", "SAKNADE ELLER OFULLSTÄNDIGA VOLYMER", "-" * 78])
    failed_volume_results = [result for result in volume_results if not result.successful]
    if failed_volume_results:
        for result in failed_volume_results:
            lines.append(
                f"- {result.relative_dir} | utdatamapp finns: {result.output_dir_exists} | "
                f"godkända PDF: {result.successful_pdfs}/{result.total_pdfs}"
            )
    else:
        lines.append("Inga avvikelser hittades.")

    lines.extend(["", "PDF-FILER UTAN NUMRERAD UTDATA", "-" * 78])
    failed_pdf_results = [result for result in pdf_results if not result.successful]
    if failed_pdf_results:
        for result in failed_pdf_results:
            relative_pdf = result.input_pdf.relative_to(input_root)
            expected = f"{result.input_pdf.stem}_1.pdf"
            lines.append(f"- {relative_pdf} | förväntat exempel: {expected}")
    else:
        lines.append("Alla ursprungliga PDF-filer har minst en numrerad utdatafil.")

    lines.extend(["", "LUCKOR I NUMRERINGEN", "-" * 78])
    gaps = [result for result in pdf_results if result.missing_numbers]
    if gaps:
        for result in gaps:
            relative_pdf = result.input_pdf.relative_to(input_root)
            numbers = ", ".join(map(str, result.missing_numbers))
            lines.append(f"- {relative_pdf} | saknade delnummer: {numbers}")
    else:
        lines.append("Inga luckor hittades i numreringen.")

    lines.extend(["", "INDEXRADER UTAN PERSONNUMMER", "-" * 78])
    if index_result.error:
        lines.append(f"Kontrollen kunde inte genomföras: {index_result.error}")
    elif index_result.missing_rows:
        lines.extend(f"- {description}" for description in index_result.missing_rows)
    else:
        lines.append("Alla indexrader har personnummer.")

    passed = failed_volumes == 0 and failed_pdfs == 0 and not index_result.error and index_result.missing == 0
    lines.extend(["", "SLUTRESULTAT", "-" * 78, "GODKÄND" if passed else "EJ GODKÄND", ""])

    report_file.parent.mkdir(parents=True, exist_ok=True)
    report_file.write_text("\n".join(lines), encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Validera examensbevisflödets volymer, delade PDF-filer och index.")
    parser.add_argument("--input", required=True, type=Path, help="Rotmapp med ursprungliga volymer och PDF-filer.")
    parser.add_argument("--output", required=True, type=Path, help="Rotmapp med bearbetade volymer och delade PDF-filer.")
    parser.add_argument("--index", type=Path, help="Indexfil i XLSX, XLSM eller CSV-format.")
    parser.add_argument("--report", type=Path, default=Path("validation_report.txt"), help="TXT-rapportens sökväg.")
    parser.add_argument("--personnummer-column", help="Exakt namn på personnummerkolumnen.")
    parser.add_argument("--sheet", help="Excel-bladets namn. Det aktiva bladet används annars.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_root = args.input.expanduser().resolve()
    output_root = args.output.expanduser().resolve()
    report_file = args.report.expanduser().resolve()
    selected_index = args.index.expanduser().resolve() if args.index else None

    if not input_root.is_dir():
        print(f"FEL: Indatamappen finns inte eller är inte en mapp: {input_root}", file=sys.stderr)
        return 1
    if not output_root.is_dir():
        print(f"FEL: Utdatamappen finns inte eller är inte en mapp: {output_root}", file=sys.stderr)
        return 1

    pdf_results = validate_pdfs(input_root, output_root)
    volume_results = validate_volumes(input_root, output_root, pdf_results)
    index_file, discovery_error = locate_index_file(output_root, selected_index)
    index_result = validate_index(index_file, discovery_error, args.personnummer_column, args.sheet)
    create_report(report_file, input_root, output_root, pdf_results, volume_results, index_result)

    successful_pdfs = sum(result.successful for result in pdf_results)
    successful_volumes = sum(result.successful for result in volume_results)
    print(f"Volymer: {successful_volumes}/{len(volume_results)} fullständigt bearbetade")
    print(f"PDF-filer: {successful_pdfs}/{len(pdf_results)} har numrerad utdata")
    if index_result.error:
        print(f"Indexkontroll: MISSLYCKADES: {index_result.error}")
    else:
        print(f"Personnummer: {index_result.extracted} extraherade, {index_result.missing} saknas")
    print(f"Rapport: {report_file}")

    has_deviations = (
        successful_volumes != len(volume_results)
        or successful_pdfs != len(pdf_results)
        or index_result.error is not None
        or index_result.missing > 0
    )
    return 2 if has_deviations else 0


if __name__ == "__main__":
    raise SystemExit(main())
